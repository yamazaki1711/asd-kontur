from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "mvp" / "FUNCTIONAL_MODEL_v0.1.md"
OUT = ROOT / "deliverables" / "Функциональная_модель_АСД-КОНТУР_MVP_Сопровождение_v0.1.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT = "F2F4F7"
BLUE_LIGHT = "E8EEF5"
GREEN_LIGHT = "E6F0EA"
GREEN = "285C45"
WHITE = "FFFFFF"


def set_font(run, name="Calibri", size=11, bold=None, italic=None, color="000000"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rpr.rFonts.set(qn(f"w:{key}"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:cantSplit")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        prevent_split(row)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_inline_runs(paragraph, text, size=11, color="000000"):
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos:match.start()]), size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), size=size, bold=True, color=color)
        else:
            set_font(paragraph.add_run(token[1:-1]), name="Consolas", size=max(size - 0.5, 8), color=DARK_BLUE)
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]), size=size, color=color)


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(paragraph.add_run("Страница "), size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def new_decimal_num_id(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_numbered_paragraph(doc, text, num_id):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    ppr.append(num_pr)
    add_inline_runs(p, text)
    return p


def add_callout(doc, label, text, fill=BLUE_LIGHT, color=INK):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_repeat_header(table.rows[0])
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(label + " "), size=10.5, bold=True, color=color)
    add_inline_runs(p, text, size=10.5, color=color)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def column_widths(headers):
    n = len(headers)
    if n == 2:
        return [2780, 6580]
    if n == 3:
        if headers[0] in ("ID", "Сущность", "Роль"):
            return [1500, 3760, 4100]
        return [3000, 3200, 3160]
    if n == 4:
        if headers[0] == "ID":
            return [900, 2950, 5510, 0][:4]
        return [2200, 1700, 1700, 3760]
    return [9360 // n] * n


def add_table(doc, headers, rows):
    widths = column_widths(headers)
    if sum(widths) != 9360:
        widths[-1] += 9360 - sum(widths)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_repeat_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell, BLUE_LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_inline_runs(p, header, size=8.8, color=INK)
        for run in p.runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, value, size=8.6)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, before, after, color in (
        (1, 16, 18, 10, BLUE),
        (2, 13, 14, 7, BLUE),
        (3, 12, 10, 5, DARK_BLUE),
    ):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(112)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("ПРОЕКТНАЯ ДОКУМЕНТАЦИЯ"), size=11, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("АСД-КОНТУР"), size=30, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Функциональная модель MVP «СОПРОВОЖДЕНИЕ»"), size=17, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(46)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Редакция 0.1"), size=12, color=MUTED)

    metadata = [
        ("Статус", "Проектная редакция для проверки на пилотном ОКС"),
        ("Дата", "11 августа 2026 года"),
        ("Владелец продукта", "Олег Щербаков"),
        ("Репозиторий", "asd-kontur"),
        ("Основание", "Паспорт проектирования АСД v0.1; ADR-0001—ADR-0004"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.65)
        p.paragraph_format.right_indent = Inches(0.65)
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run(label + ": "), size=10.5, bold=True, color=INK)
        set_font(p.add_run(value), size=10.5)

    add_callout(
        doc,
        "Назначение.",
        "Задать проверяемое функциональное поведение первого MVP: от актуальной РД и партии МТР до полевого факта, проекта ИД и готовности объёма к предъявлению.",
        fill=GREEN_LIGHT,
        color=GREEN,
    )
    doc.add_page_break()


def add_front_matter(doc):
    p = doc.add_paragraph(style="Heading 1")
    add_inline_runs(p, "Управленческое резюме", size=16, color=BLUE)
    for text in (
        "Функциональная модель фиксирует узкий, но сквозной MVP режима «СОПРОВОЖДЕНИЕ» на одном пилотном ОКС. Система должна не просто хранить файлы, а заранее показывать, какие доказательства обязательны, связывать их с работой, МТР, контролем и объёмом и не позволять скрыть критический дефицит.",
        "Базовый принцип доверия: детерминированное ядро рассчитывает применимость, блокировки и статусы; ИИ создаёт проверяемые кандидаты и проекты; уполномоченный человек подтверждает факты и решения. Любой значимый вывод воспроизводим по версии источника, локатору и правилу.",
        "Полевой контур проектируется offline-first. Повторная синхронизация не создаёт дублей, конфликты не разрешаются молча, а подтверждённая история не переписывается.",
    ):
        p = doc.add_paragraph(style="Normal")
        add_inline_runs(p, text)
    add_callout(doc, "Критерий ценности.", "Дефицит должен быть обнаружен до сокрытия или продолжения работ, когда доказательство ещё можно получить законно и технически.")

    p = doc.add_paragraph(style="Heading 1")
    add_inline_runs(p, "Содержание", size=16, color=BLUE)
    titles = [
        "1. Назначение документа", "2. Цель MVP и проверяемая гипотеза", "3. Границы MVP",
        "4. Принципы функционального поведения", "5. Роли и ответственность", "6. Модель доступа",
        "7. Основные сущности", "8. Обязательные инварианты", "9. Жизненные циклы",
        "10. Сквозной основной сценарий", "11. Исключительные сценарии", "12. Правила проверок и реакций",
        "13. Граница детерминированного ядра, ИИ и человека", "14. Функциональные контуры и экраны",
        "15. Полевой RAG-клиент", "16. Уведомления и задания", "17. Отчёты и выходные результаты MVP",
        "18. Нефункциональные требования, влияющие на функцию", "19. Критерии приёмки MVP",
        "20. Пилотные данные и сценарий демонстрации", "21. Метрики пилота", "22. Открытые решения",
        "23. Следующие артефакты", "24. Условие утверждения редакции",
    ]
    for title in titles:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2)
        add_inline_runs(p, title, size=9.5)


def parse_markdown_body(doc, text):
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1. "))
    lines = lines[start:]
    i = 0
    current_num_id = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            headers = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, headers, rows)
            current_num_id = None
            continue
        if line.startswith("### "):
            current_num_id = None
            p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, line[4:], size=13, color=BLUE)
        elif line.startswith("## "):
            current_num_id = None
            if line.startswith("## 12. "):
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            add_inline_runs(p, line[3:], size=16, color=BLUE)
        elif line.startswith("> "):
            current_num_id = None
            add_callout(doc, "Сквозная цепочка.", line[2:])
        elif line.startswith("- "):
            current_num_id = None
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            if current_num_id is None:
                current_num_id = new_decimal_num_id(doc)
            add_numbered_paragraph(doc, re.sub(r"^\d+\. ", "", line), current_num_id)
        else:
            current_num_id = None
            p = doc.add_paragraph(style="Normal")
            add_inline_runs(p, line)
        i += 1


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.48)
    section.footer_distance = Inches(0.48)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run("АСД-КОНТУР  |  MVP «СОПРОВОЖДЕНИЕ»  |  Функциональная модель v0.1"), size=9, color=MUTED)
    set_page_number(section.footer.paragraphs[0])

    add_cover(doc)
    add_front_matter(doc)
    parse_markdown_body(doc, SOURCE.read_text(encoding="utf-8"))

    core = doc.core_properties
    core.title = "АСД-КОНТУР — Функциональная модель MVP «СОПРОВОЖДЕНИЕ» v0.1"
    core.subject = "Функциональная модель первого MVP комплекса АСД-КОНТУР"
    core.author = "Олег Щербаков"
    core.keywords = "АСД-КОНТУР, строительство, исполнительная документация, MVP, сопровождение"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
